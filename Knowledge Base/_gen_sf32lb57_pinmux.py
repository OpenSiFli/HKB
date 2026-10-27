# -*- coding: utf-8 -*-
"""
Generate wiki/mcu/all/pinmux/SF32LB57_Pinmux.md from the official pinmux xlsx.
Mirrors the canonical SF32LB52_Pinmux.md format.
"""
import openpyxl, sys, os
sys.stdout.reconfigure(encoding='utf-8')

SRC = r'E:/Work/Package/ButterFli-Micro-S/pinmux/pinmux - 202600901.xlsx'
OUT = r'D:/obsidian/硬件知识库/wiki/mcu/all/pinmux/SF32LB57_Pinmux.md'
RAW = 'raw/mcu/all_chip/pinmux/pinmux - 202600901.xlsx'

wb = openpyxl.load_workbook(SRC, data_only=True)

# ---------- HPSYS ----------
ws = wb['HPSYS']
func_cols = list(range(8, 24))  # Function 0..6, 8..16
rows = []
for r in range(2, ws.max_row + 1):
    pad = ws.cell(row=r, column=3).value
    if pad is None:
        continue
    pad = str(pad).strip()
    q68 = ws.cell(row=r, column=4).value
    q80 = ws.cell(row=r, column=5).value
    ptype = ws.cell(row=r, column=6).value
    pupd = ws.cell(row=r, column=7).value
    funcs = {}
    for c in func_cols:
        v = ws.cell(row=r, column=c).value
        if v is not None and str(v).strip() != '':
            lbl = str(ws.cell(row=1, column=c).value).replace(' ', '').replace('Function', 'F')
            funcs[lbl] = str(v).replace('\n', ' ').strip()
    rows.append(dict(pad=pad, q68=q68, q80=q80, ptype=ptype, pupd=pupd, funcs=funcs))

def yn(v):
    return 'Y' if v not in (None, '') else '—'

def fn_num(k):
    return int(k[1:])

big_rows = []
for x in rows:
    pad = x['pad'].replace('PAD_', '')
    default = x['funcs'].get('F0', '')
    dedicated = []
    special = []
    for k, v in x['funcs'].items():
        if k == 'F0':
            continue
        if v.startswith('MATRIX_DIO'):
            continue  # crossbar, noted globally
        if v.startswith('#'):
            special.append('%s(%s)' % (v, k))
        else:
            dedicated.append('%s(%s)' % (v, k))
    dedicated.sort(key=lambda s: fn_num(s.split('(')[1].rstrip(')')))
    special.sort(key=lambda s: fn_num(s.split('(')[1].rstrip(')')))
    big_rows.append('| %s | %s | %s | %s | %s | %s | %s | %s |' % (
        pad, yn(x['q68']), yn(x['q80']),
        (x['pupd'] or ''), (x['ptype'] or ''),
        default,
        (' / '.join(dedicated) if dedicated else '—'),
        (' / '.join(special) if special else '—'),
    ))

# ---------- matrix index ----------
ws2 = wb['matrix index']
mod_idx = {}  # module -> list of (idx, func)
cur_mod = None
for r in range(2, ws2.max_row + 1):
    m = ws2.cell(row=r, column=1).value
    idx = ws2.cell(row=r, column=3).value
    func = ws2.cell(row=r, column=5).value
    if m and str(m).strip():
        cur_mod = str(m).strip()
    if idx is None or func is None:
        continue
    mod_idx.setdefault(cur_mod, []).append((int(idx), str(func).strip()))

matrix_tables = []
for mod in ['I2C', 'SPI', 'UART', 'CAN', 'PTA/BT', 'GPTIM', 'LPTIM', 'ATIM', 'PWM', 'I2S', 'PTM', 'LDCD', 'PDM', 'DCMI']:
    lst = mod_idx.get(mod, [])
    if not lst:
        continue
    lines = ['### %s（%d 个信号）' % (mod, len(lst)), '', '| idx | func |', '|---|---|']
    for idx, func in lst:
        lines.append('| %d | %s |' % (idx, func))
    matrix_tables.append('\n'.join(lines))

# ---------- lcdc mux ----------
ws3 = wb['lcdc mux']
lcdc_rows = []
for r in range(4, ws3.max_row + 1):  # data starts at row 4 (lcdc_if[0])
    net = ws3.cell(row=r, column=2).value
    spi = ws3.cell(row=r, column=3).value
    jdi = ws3.cell(row=r, column=4).value
    epd = ws3.cell(row=r, column=5).value
    dbi = ws3.cell(row=r, column=6).value
    dpi = ws3.cell(row=r, column=7).value
    cnt = ws3.cell(row=r, column=8).value
    if net is None:
        continue
    def c(v):
        return '' if v is None else str(v).replace('\n', ' ').strip()
    lcdc_rows.append('| %s | %s | %s | %s | %s | %s | %s |' % (c(net), c(spi), c(jdi), c(epd), c(dbi), c(dpi), c(cnt)))

ws4 = wb['lcdc dpi inner mux']
dpi_rows = []
for r in range(2, ws4.max_row + 1):
    net = ws4.cell(row=r, column=1).value
    mode = ws4.cell(row=r, column=2).value
    if net is None:
        continue
    def c(v):
        return '' if v is None else str(v).replace('\n', ' ').strip()
    cols = [c(ws4.cell(row=r, column=cc).value) for cc in range(3, 11)]  # mode 0..7
    dpi_rows.append('| %s | %s | %s |' % (c(net), c(mode), ' | '.join(cols)))

n_q68 = sum(1 for x in rows if yn(x['q68']) == 'Y')
n_total = len(rows)

# ---------- compose ----------
L = []
L.append('---')
L.append('title: SF32LB57 引脚复用矩阵（Pinmux）')
L.append('type: mcu_pinmux')
L.append('aliases: ["SF32LB57_Pinmux","SF32LB57 引脚复用","SF32LB57 pinmux"]')
L.append('tags: ["企业MCU","引脚复用/Pinmux","芯片系列-SF32LB57","原理图审图/故障分析/报告输出"]')
L.append('sources: [[SF32LB57X_Datasheet_V0.2_20260909],[SF32LB57x_batt_Datasheet_V0.2_20260909]]')
L.append('mcu_family: ["SF32LB57"]')
L.append('series_applicability: ["通用-两系列共用"]')
L.append('mcu_model: ["SF32LB57X","SF32LB57x"]')
L.append('hardware_version: ["pinmux_202600901","QFN68 7x7","QFN80 8x8","BGA112 5.4x4.4","Datasheet V0.2"]')
L.append('created: 2026-09-09')
L.append('updated: 2026-09-09')
L.append('---')
L.append('')
L.append('# SF32LB57 家族 PINMUX（引脚复用矩阵）')
L.append('')
L.append('> **权威来源**：`pinmux - 202600901.xlsx`（已归档 `raw/mcu/all_chip/pinmux/`）。')
L.append('> **规则**：PINMUX 相关内容**以此表格为准**；规格书（DS5702 3.3V / DS5701 电池）Table 5-2/5-4 若与下表有差异，原则上以本表为准，差异在文末「与规格书差异」一节标注。')
L.append('> **系列适用性**：本表覆盖 GPIO Group A（PA00–PA57），**电池供电系列（SF32LB57x）与 3.3V 供电系列（SF32LB57X）共用同一套大核 GPIO 复用矩阵**；两系列差异仅在于专用电源/模拟管脚（见 [[SF32LB57_Peripheral_Power]] §1），不涉及 GPIO Group A 的复用分配。')
L.append('')
L.append('## 1. 工作表（Sheet）结构与含义')
L.append('')
L.append('| 工作表 | 含义 | 是否入库 |')
L.append('|---|---|---|')
L.append('| **HPSYS** | 大核（HCPU / Big Core）GPIO 复用矩阵（GPIO Group A，PA00–PA57） | ✅ 入库（见第 3 节） |')
L.append('| **matrix index** | 矩阵交叉开关信号总表（160 个可路由信号，idx 16–175） | ✅ 入库（见第 4 节） |')
L.append('| **lcdc mux** | LCD 控制器接口复用（IF_SEL 0–3：SPI/JDI/EPD/DBI/DPI） | ✅ 入库（见第 5 节） |')
L.append('| **lcdc dpi inner mux** | DPI 内部 RGB 位宽复用（mode 0–7） | ✅ 入库（见第 5 节） |')
L.append('| **simplified** | 精简视图（PAD + PUPD + Func 0/1/2/3/4/5/6/15） | 仅对照，不另列表（内容与 HPSYS 一致） |')
L.append('')
L.append('> ⚠️ **本文件 pinmux xlsx 仅含 HPSYS（大核 GPIO Group A）**，未含 LPSYS（小核 GPIO）/ PBR（电源 RF 控制）工作表（SF32LB52 家族 pinmux 含 LPSYS/PBR，57 家族本表未提供，见「与规格书差异」第 3 条）。')
L.append('')
L.append('## 2. PINMUX 类型说明（命名规则）')
L.append('')
L.append('- **所有 PAxx 默认功能 = `GPIO_Ax`（Func0）**：未做复用配置时即为通用 GPIO。')
L.append('- **`MATRIX_DIOx`（Func16，每个 PA 均有）**：矩阵交叉开关输出。SF32LB57 采用**灵活矩阵（matrix）架构**——任一矩阵信号（见第 4 节：I2C/UART/TIM/PWM/… 共 160 个）均可经 `MATRIX_DIOx` 路由到**任意 PA 引脚**，替代 SF32LB52 家族「`PAxx_I2C_UART` / `PAxx_TIM` 固定通用分配组」的写法。本表不在每行重复列出 `MATRIX_DIOx`。')
L.append('- **专用复用（Func1/2/3/…）**：该 IO 上固化的特定外设功能（如 SPI/I2S/MPI/SD/LCDC/USART/SWD/USB 等），见第 3 节「专用复用」列。')
L.append('- **特殊标记前缀 `#`**：模拟 / 测试 / 唤醒 / 时钟类专用功能（`#GPADC_CHx` ADC 通道、`#WKUP_PINx` 唤醒、`#XTAL32K_XI/XO/EXT` 32K 晶振、`#USB11_DP/DM` USB、`#EFUSE_PWR` 熔丝电源、`#RTC_CLK`/`#LPTIM2_OUT`/`#LPTIM2_OUTB` 时钟/定时输出、`#ADC1N`/`#ADC2P`/`#ADC2N` 音频 ADC）。')
L.append('')
L.append('## 3. 大核 GPIO（HPSYS / GPIO Group A）PINMUX 总表')
L.append('')
L.append('> 列说明：**QFN68** = 7×7 封装是否引出（Y/—）；**QFN80/BGA112** = 8×8 / 5.4×4.4 封装是否引出；**PUPD** = 上下拉默认；**PadType** = 管脚电气类型（TYPE1/TYPE2）；**默认 Func0** = 上电默认 GPIO；**专用复用** = 固化外设功能（括号标 Func 选择器号）；**特殊标记** = `#` 类模拟/唤醒/时钟功能。')
L.append('> **物理引脚号（Pin#）本表未列**，取自规格书 Table 5-2/5-3（见「与规格书差异」第 1 条）。所有 PA 均含 `MATRIX_DIOx`（Func16，矩阵交叉开关）未逐行列出。')
L.append('')
L.append('| PA | QFN68 | QFN80/BGA112 | PUPD | PadType | 默认 Func0 | 专用复用 (FuncN) | 特殊标记 (#) |')
L.append('|---|---|---|---|---|---|---|---|')
L.extend(big_rows)
L.append('')
L.append('> 注：QFN68 封装未引出的 PA = **PA21、PA45、PA46、PA47、PA48、PA49、PA50、PA51、PA52、PA53、PA54、PA57**（共 %d 个 PA 仅存在于 QFN80/BGA112）；故 QFN68 = %d GPIO，QFN80/BGA112 = %d GPIO。' % (len(rows) - n_q68, n_q68, n_total))
L.append('')
L.append('## 4. Matrix Index（矩阵交叉开关信号表，160 个）')
L.append('')
L.append('> 下表为 `matrix index` 工作表，列出矩阵可路由的 160 个信号（idx 16–175）。任意信号可经 `MATRIX_DIOx`（Func16）路由到任意 PA 引脚（见第 2 节）。按模块分组：')
L.append('')
L.extend(matrix_tables)
L.append('')
L.append('## 5. LCD 接口复用（LCDC）')
L.append('')
L.append('### 5.1 lcdc mux（IF_SEL 选择接口类型）')
L.append('')
L.append('> 括号数字（8 / 17 / 22 / 24 / 30）为各接口的 IF_SEL 子接口选择器索引；末列数字为该 lcdc_if 信号所占 PAD 数（4/3/2）。`lcdc_if[0..29]` 共 30 个 LCD 接口信号。')
L.append('')
L.append('| lcdc_if | IF_SEL=0 SPI | IF_SEL=1 JDI | IF_SEL=2 EPD | IF_SEL=3 DBI | DPI | PAD数 |')
L.append('|---|---|---|---|---|---|---|')
L.extend(lcdc_rows)
L.append('')
L.append('### 5.2 lcdc dpi inner mux（DPI RGB 位宽复用）')
L.append('')
L.append('> mode 0–7 选择 RGB 数据位宽与映射（rgb565/rgb666/rgb888 及其 over-8bit 变体）；software default = mode 1。')
L.append('')
L.append('| matrix net | dpi_d | 0(rgb565 conf1) | 1(rgb565 conf2) 默认 | 2(rgb565 conf3) | 3(rgb666 conf1) | 4(rgb666 conf2) | 5(rgb888) | 6(rgb666 over8) | 7(rgb888 over8) |')
L.append('|---|---|---|---|---|---|---|---|---|---|')
L.extend(dpi_rows)
L.append('')
L.append('## 6. 与规格书（DS5702 3.3V / DS5701 电池）差异 / 待核实')
L.append('')
L.append('1. **物理引脚号（Pin#）本表缺失**：pinmux xlsx 仅给 QFN68 / QFN80-BGA112 的 Y/N 引出标记，未列具体封装引脚号；物理 Pin# 须查规格书 Table 5-2（大核 GPIO）/ Table 5-3（专用管脚）。待补。')
L.append('2. **GPIO 数量**：本表 QFN68 = %d GPIO（PA21、PA45–PA54、PA57 未引出）；规格书页 [[SF32LB57]] 标注 QFN68 = 47 GPIO，二者存在偏差，**以本 pinmux 表为准**，规格书 GPIO 计数待原厂复核（草稿 V0.2）。' % n_q68)
L.append('3. **LPSYS / 小核 GPIO 缺失**：本 pinmux xlsx 仅含 HPSYS（大核 GPIO Group A），无 LPSYS/PBR 工作表；SF32LB57 为双大核架构，小核 GPIO 复用是否另行提供待用户手册确认。')
L.append('4. **MATRIX_DIOx 矩阵路由**：57 家族以矩阵交叉开关实现 I2C/UART/TIM/PWM 等通用分配（替代 52 家族固定 `PAxx_I2C_UART`/`PAxx_TIM` 列），具体矩阵信号见第 4 节。')
L.append('5. **专用复用外设命名差异（相对 52 家族）**：57 为 **MPI3**（非 MPI2）、**SD1/SD2**（双 SD 控制器）、**USART1**（单 Debug UART，PA18/PA19）、**无独立 PA_I2C_UART/PA_TIM 列**；禁止跨家族套用 52 家族引脚分配。')
L.append('')
L.append('## 关联页面')
L.append('- 家族：[[SF32LB57]]｜3.3V：[[SF32LB57X]]｜电池：[[SF32LB57x_batt|SF32LB57x]]')
L.append('- GPIO：[[SF32LB57_Peripheral_GPIO]]｜电源：[[SF32LB57_Peripheral_Power]]｜存储：[[SF32LB57_Peripheral_Memory]]｜串口：[[SF32LB57_Peripheral_Serial]]')
L.append('- 52 家族对照：[[SF32LB52_Pinmux]]')
L.append('- 来源：[[SF32LB57X_Datasheet_V0.2_20260909]]、[[SF32LB57x_batt_Datasheet_V0.2_20260909]]')
L.append('')
L.append('## 原始来源完整引用')
L.append('1. [pinmux - 202600901.xlsx]，Raw：`%s`（原件 `E:/Work/Package/ButterFli-Micro-S/pinmux/pinmux - 202600901.xlsx` 仅复制未移动）。' % RAW)
L.append('2. [DS5702-SF32LB57X V0.2]，Raw：`raw/mcu/all_chip/datasheets/DS5702-SF32LB57X-芯片技术规格书-V0p2-20260909.pdf`，Table 5-2（GPIO 复用）、Table 5-4（Matrix Index）。')
L.append('3. [DS5701-SF32LB57x V0.2]，Raw：`raw/mcu/all_chip/datasheets/DS5701-SF32LB57x-芯片技术规格书-V0p2-20260909.pdf`。')

txt = '\n'.join(L)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, 'w', encoding='utf-8') as f:
    f.write(txt + '\n')
print('WROTE', OUT, 'bytes=', len(txt.encode('utf-8')))
print('HPSYS pads:', n_total, '| QFN68 GPIO:', n_q68)
print('matrix modules:', len(matrix_tables))
